"""Excel and PDF export of a calculated voyage (MVP report layout)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

from .fonts import ensure_fonts

FONT, FONT_BOLD = ensure_fonts()

GRID_COLUMNS = [
    ("tank_id", "Tank"),
    ("parcel_id", "Parcel"),
    ("ullage", "Ullage (cm)"),
    ("trim_correction", "Trim Corr"),
    ("corrected_ullage", "Corr. Ullage"),
    ("tov", "TOV (m³)"),
    ("therm_corr", "Therm."),
    ("gov", "GOV (m³)"),
    ("fill_percent", "Fill %"),
    ("temp_celsius", "Temp °C"),
    ("density_vac", "Dens (Vac)"),
    ("density_air", "Dens (Air)"),
    ("vcf", "VCF"),
    ("gsv", "GSV (m³)"),
    ("mt_air", "MT (Air)"),
    ("mt_vac", "MT (Vac)"),
]

FORMATS = {
    "ullage": "{:.1f}", "trim_correction": "{:.1f}", "corrected_ullage": "{:.1f}",
    "tov": "{:.3f}", "therm_corr": "{:.6f}", "gov": "{:.3f}", "fill_percent": "{:.1f}",
    "temp_celsius": "{:.1f}", "density_vac": "{:.4f}", "density_air": "{:.4f}",
    "vcf": "{:.5f}", "gsv": "{:.3f}", "mt_air": "{:.3f}", "mt_vac": "{:.3f}",
}


def _fmt(key: str, value) -> str:
    if value is None:
        return ""
    if key in FORMATS and isinstance(value, (int, float)):
        return FORMATS[key].format(value)
    return str(value)


def _header_lines(ship_name: str, voyage: dict, totals: dict) -> list[tuple[str, str]]:
    return [
        ("Ship", ship_name),
        ("Voyage No", voyage.get("voyage_number", "")),
        ("Date", voyage.get("date", "")),
        ("Port / Terminal", f"{voyage.get('port', '')} / {voyage.get('terminal', '')}"),
        ("V.E.F.", f"{voyage.get('vef', 1.0):.5f}"),
        ("Draft AFT / FWD", f"{voyage.get('draft_aft', 0.0):.2f} / {voyage.get('draft_fwd', 0.0):.2f}"),
        ("Trim", f"{totals.get('trim', 0.0):.2f}"),
        ("Chief Officer", voyage.get("chief_officer", "")),
        ("Master", voyage.get("master", "")),
    ]


def export_xlsx(ship_name: str, voyage: dict, calc: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ullage Report"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "ULLAGE REPORT"
    ws["A1"].font = Font(bold=True, size=14)

    row = 3
    for label, value in _header_lines(ship_name, voyage, calc["totals"]):
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    header_row = row
    for col, (_, title) in enumerate(GRID_COLUMNS, start=1):
        c = ws.cell(row=header_row, column=col, value=title)
        c.font = bold
        c.fill = header_fill
        c.border = border
        c.alignment = Alignment(horizontal="center")

    for r in calc["rows"]:
        row += 1
        for col, (key, _) in enumerate(GRID_COLUMNS, start=1):
            value = r.get(key)
            c = ws.cell(row=row, column=col)
            if isinstance(value, (int, float)) and key not in ("tank_id", "parcel_id"):
                c.value = float(_fmt(key, value)) if _fmt(key, value) else None
            else:
                c.value = value if value is not None else ""
            c.border = border

    row += 1
    ws.cell(row=row, column=1, value="TOTALS").font = bold
    gsv_col = next(i for i, (k, _) in enumerate(GRID_COLUMNS, start=1) if k == "gsv")
    mt_col = next(i for i, (k, _) in enumerate(GRID_COLUMNS, start=1) if k == "mt_air")
    ws.cell(row=row, column=gsv_col, value=round(calc["totals"]["gsv"], 3)).font = bold
    ws.cell(row=row, column=mt_col, value=round(calc["totals"]["mt_air"], 3)).font = bold

    if calc.get("parcels"):
        row += 2
        ws.cell(row=row, column=1, value="PARCEL SUMMARY").font = bold
        row += 1
        titles = ["Parcel", "Grade", "Receiver", "GSV", "MT (Air)", "Ship w/ VEF", "B/L", "Diff", "Diff ‰"]
        for col, t in enumerate(titles, start=1):
            c = ws.cell(row=row, column=col, value=t)
            c.font = bold
            c.fill = header_fill
        for p in calc["parcels"]:
            row += 1
            values = [
                p["parcel_id"], p["name"], p["receiver"],
                round(p["gsv"], 3), round(p["mt_air"], 3), round(p["ship_with_vef"], 3),
                round(p["bl_figure"], 3), round(p["diff_with_vef"], 3),
                round(p["diff_permille_with_vef"], 3),
            ]
            for col, v in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=v)

    for col in range(1, len(GRID_COLUMNS) + 1):
        ws.column_dimensions[ws.cell(row=header_row, column=col).column_letter].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(ship_name: str, voyage: dict, calc: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    styles["Title"].fontName = FONT_BOLD
    styles["Heading3"].fontName = FONT_BOLD
    story = [Paragraph(f"ULLAGE REPORT — {ship_name}", styles["Title"])]

    header_data = [[label, value] for label, value in _header_lines(ship_name, voyage, calc["totals"])]
    header_table = Table(header_data, colWidths=[45 * mm, 90 * mm], hAlign="LEFT")
    header_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTNAME", (0, 0), (0, -1), FONT_BOLD),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
    ]))
    story += [header_table, Spacer(1, 6)]

    data = [[title for _, title in GRID_COLUMNS]]
    for r in calc["rows"]:
        data.append([_fmt(key, r.get(key)) for key, _ in GRID_COLUMNS])
    totals_row = [""] * len(GRID_COLUMNS)
    totals_row[0] = "TOTALS"
    gsv_i = next(i for i, (k, _) in enumerate(GRID_COLUMNS) if k == "gsv")
    mt_i = next(i for i, (k, _) in enumerate(GRID_COLUMNS) if k == "mt_air")
    totals_row[gsv_i] = f"{calc['totals']['gsv']:.3f}"
    totals_row[mt_i] = f"{calc['totals']['mt_air']:.3f}"
    data.append(totals_row)

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ("FONTNAME", (0, -1), (-1, -1), FONT_BOLD),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F5F7FB")]),
    ]))
    story.append(table)

    if calc.get("parcels"):
        story.append(Spacer(1, 8))
        story.append(Paragraph("Parcel Summary", styles["Heading3"]))
        pdata = [["Parcel", "Grade", "Receiver", "GSV", "MT (Air)", "Ship w/ VEF", "B/L", "Diff", "Diff ‰"]]
        for p in calc["parcels"]:
            pdata.append([
                p["parcel_id"], p["name"], p["receiver"],
                f"{p['gsv']:.3f}", f"{p['mt_air']:.3f}", f"{p['ship_with_vef']:.3f}",
                f"{p['bl_figure']:.3f}", f"{p['diff_with_vef']:.3f}",
                f"{p['diff_permille_with_vef']:.3f}",
            ])
        ptable = Table(pdata, hAlign="LEFT")
        ptable.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("FONTNAME", (0, 0), (-1, -1), FONT),
            ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
        ]))
        story.append(ptable)

    doc.build(story)
    return buf.getvalue()
